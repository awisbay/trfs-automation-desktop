import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook,load_workbook
from audit.radio_sharing_audit import audit_radio_sharing
from audit.dump_parser import parse_dump
from test_audit_termpoint import cmdump

NODE='MIN823_SITEB01'
OTHER='MIN823_SITEB02'


def fixture(node=NODE,flag='false',serial='ABC'):
    return {
        f'ManagedElement={node},Equipment=1,FieldReplaceableUnit=BB-1':{},
        f'ManagedElement={node},Equipment=1,FieldReplaceableUnit=B28_RRU1':
            {'serialNumber':serial,'isSharedWithExternalNE':flag},
        f'ManagedElement={node},RiLink=1':{
            'riPortRef1':'Equipment=1,FieldReplaceableUnit=BB-1,RiPort=A',
            'riPortRef2':'Equipment=1,FieldReplaceableUnit=B28_RRU1,RiPort=DATA_1'}}


def lld(path,shared='No',second=None):
    wb=Workbook()
    ws=wb.active
    ws.title='CPRI connectivity'
    ws.append(['Notes'])
    ws.append(['PLA ID','BBID','BB RI Port','Radio DATA Port','Radio Shared between BB'])
    ws.append(['MIN823','BB1','A','DATA_1',shared])
    if second is not None:
        ws.append(['MIN823','BB2','A','DATA_1',second])
    wb.save(path)
    wb.close()


class SharingTests(unittest.TestCase):
    def test_flag_on_radio_riport_uses_actual_mo(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'lld.xlsx'
            lld(path)
            records=fixture()
            fru=f'ManagedElement={NODE},Equipment=1,FieldReplaceableUnit=B28_RRU1'
            del records[fru]['isSharedWithExternalNE']
            records[fru+',RiPort=DATA_1']={'isSharedWithExternalNE':'true'}
            rows=audit_radio_sharing(records,lld_path=path)
            self.assertEqual((rows[0].mo,rows[0].status),
                             ('Equipment=1,FieldReplaceableUnit=B28_RRU1,RiPort=DATA_1','Mismatch'))
            records[fru+',RiPort=DATA_2']={'isSharedWithExternalNE':'false'}
            rows=audit_radio_sharing(records,lld_path=path)
            self.assertEqual([r.status for r in rows],['Mismatch','NotFound'])

    def test_nonshared_and_wrong_true_both_formats(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'lld.xlsx'
            lld(path)
            for flag,status in [('false','Match'),('true','Mismatch')]:
                records=fixture(flag=flag)
                dcg='\n'.join('MO  '+dn+'\n'+'\n'.join(f'{k}  {v}' for k,v in a.items()) for dn,a in records.items()).encode()
                for kind,payload in [('xml',cmdump(records)),('log',dcg)]:
                    dump=Path(folder)/('dump.'+kind)
                    dump.write_bytes(payload)
                    rows=audit_radio_sharing(parse_dump(str(dump)),lld_path=path)
                    self.assertEqual(len(rows),1)
                    self.assertEqual((rows[0].expected,rows[0].status),('false',status))
                    self.assertIn('RI A',rows[0].remark)

    def test_shared_same_serial_and_selection(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'lld.xlsx'
            lld(path,'Yes','Yes')
            records=fixture(flag='false')
            records.update(fixture(OTHER,'true'))
            rows=audit_radio_sharing(records,lld_path=path)
            self.assertEqual([(r.expected,r.status) for r in rows],[('true','Mismatch'),('true','Match')])
            rows=audit_radio_sharing(records,nodes=[NODE],lld_path=path)
            self.assertEqual((len(rows),rows[0].status),(1,'Mismatch'))
            rows=audit_radio_sharing(records)
            self.assertEqual(rows[0].expected,'true')

    def test_same_port_and_fru_name_do_not_prove_shared(self):
        records=fixture()
        records.update(fixture(OTHER,serial='DIFFERENT'))
        rows=audit_radio_sharing(records)
        self.assertTrue(all(r.status=='NotFound' for r in rows))

    def test_sync_supports_plan_and_conflict_is_unresolved(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'lld.xlsx'
            records=fixture(flag='true')
            records[f'ManagedElement={NODE},NodeGroupSyncMember=1']={
                'syncRiPortCandidate':'Equipment=1,FieldReplaceableUnit=BB-1,RiPort=A'}
            lld(path,'Yes')
            self.assertEqual(audit_radio_sharing(records,lld_path=path)[0].status,'Match')
            lld(path,'No')
            row=audit_radio_sharing(records,lld_path=path)[0]
            self.assertEqual(row.status,'NotFound')
            self.assertIn('LLD says not shared',row.remark)

    def test_incomplete_or_conflicting_evidence_never_passes(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'lld.xlsx'
            for shared in ('','invalid','Yes'):
                lld(path,shared)
                self.assertEqual(audit_radio_sharing(fixture(),lld_path=path)[0].status,'NotFound')
            lld(path)
            for mutation in ('flag','reference','port','duplicate'):
                records=fixture()
                if mutation=='flag':
                    del records[f'ManagedElement={NODE},Equipment=1,FieldReplaceableUnit=B28_RRU1']['isSharedWithExternalNE']
                elif mutation=='reference':
                    records[f'ManagedElement={NODE},RiLink=1']['riPortRef2']='ManagedElement=OTHER,FieldReplaceableUnit=B28_RRU1,RiPort=DATA_1'
                elif mutation=='port':
                    records[f'ManagedElement={NODE},RiLink=1']['riPortRef2']='Equipment=1,FieldReplaceableUnit=B28_RRU1,RiPort=DATA_2'
                else:
                    records[f'ManagedElement={NODE},RiLink=2']=dict(records[f'ManagedElement={NODE},RiLink=1'])
                rows=audit_radio_sharing(records,lld_path=path)
                self.assertEqual(rows[0].status,'NotFound',mutation)

    def test_excel_and_report_only(self):
        from audit.audit_core import write_excel,generate_moshell_scripts,generate_cmedit_scripts
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'lld.xlsx'
            lld(path)
            rows=audit_radio_sharing(fixture(flag='true'),lld_path=path)
            report=Path(folder)/'report.xlsx'
            write_excel(rows,str(report),{})
            wb=load_workbook(report)
            self.assertEqual(wb['Detail']['E2'].value,'isSharedWithExternalNE')
            self.assertIn('RI A',wb['Detail']['J2'].value)
            wb.close()
            self.assertEqual(generate_moshell_scripts(rows,folder,'SITE',str(report)),[])
            self.assertEqual(generate_cmedit_scripts(rows,folder,'SITE',str(report)),[])

if __name__=='__main__':
    unittest.main()
