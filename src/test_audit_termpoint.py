"""Offline regression tests, including both dump formats through parse_dump."""
import gzip
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from audit.dump_parser import parse_dump
from audit.termpoint_audit import audit_termpoint_to_gnb


NR = "MIN3117_SITEB01"
LTE = "MIN3117_SITEB02"
ADDRESS = "Transport=1,Router=NR,InterfaceIPv4=X2_CP_UP,AddressIPv4=1"
PARENT = "ENodeBFunction=1,GUtraNetwork=1,ExternalGNodeBFunction=" + NR
TERM = PARENT + ",TermPointToGNB=" + NR


def fixture():
    return {
        f"ManagedElement={NR},GNBDUFunction=1": {"gNBId": "902007"},
        f"ManagedElement={NR},{ADDRESS}": {"address": "10.40.128.0/31"},
        f"ManagedElement={NR},{TERM}": {"ipAddress": "10.40.128.0"},
        f"ManagedElement={LTE},ENodeBFunction=1": {"eNBId": "1"},
        f"ManagedElement={LTE},{PARENT}": {"gNodeBId": "902007"},
        f"ManagedElement={LTE},{TERM}": {"ipAddress": "0.0.0.0"},
    }


def cmdump(records):
    root = ET.Element("bulkCmConfigDataFile")
    containers = {}
    ns = "{genericNrm.xsd}"
    for dn, attrs in records.items():
        parent = root
        path = []
        for part in dn.split(","):
            path.append(part)
            key = tuple(path)
            cls, name = part.split("=", 1)
            if key not in containers:
                el = ET.SubElement(parent, ns + "VsDataContainer", id=name)
                block = ET.SubElement(el, ns + "attributes")
                ET.SubElement(block, ns + "vsDataType").text = "vsData" + cls
                payload = ET.SubElement(block, "{EricssonSpecificAttributes.xsd}vsData" + cls)
                containers[key] = (el, payload)
            parent, payload = containers[key]
        for attr, value in attrs.items():
            ET.SubElement(payload, attr).text = value
    return ET.tostring(root)


class TermPointAuditTests(unittest.TestCase):
    def test_both_dump_formats_and_archives(self):
        records = fixture()
        dcg = "\n".join("MO  " + dn + "\n" + "\n".join(
            f"{k}  {v}" for k, v in attrs.items()) for dn, attrs in records.items()).encode()
        with tempfile.TemporaryDirectory() as folder:
            for kind, payload in (("cmdump", cmdump(records)), ("modump", dcg)):
                for zipped in (False, True):
                    with self.subTest(kind=kind, zipped=zipped):
                        path = Path(folder) / (kind + (".zip" if zipped else ".txt"))
                        if zipped:
                            with zipfile.ZipFile(path, "w") as archive:
                                archive.writestr("export.xml" if kind == "cmdump" else "node_dcg_k.log.gz",
                                                 payload if kind == "cmdump" else gzip.compress(payload))
                        else:
                            path.write_bytes(payload)
                        rows = audit_termpoint_to_gnb(parse_dump(str(path)))
                        self.assertEqual([(r.node, r.status, r.expected) for r in rows],
                                         [(NR, "Match", "10.40.128.0"),
                                          (LTE, "Mismatch", "10.40.128.0")])

    def test_empty_wrong_and_matching_ip(self):
        for ip, expected in (("", "Mismatch"), ("0.0.0.0", "Mismatch"),
                             ("10.40.128.6", "Mismatch"), ("10.40.128.0", "Match")):
            with self.subTest(ip=ip):
                records = fixture()
                records[f"ManagedElement={LTE},{TERM}"]["ipAddress"] = ip
                row, = audit_termpoint_to_gnb(records, nodes=[LTE])
                self.assertEqual(row.status, expected)
                self.assertEqual(row.ref_cell, NR)

    def test_numeric_external_id_uses_parent_gnb_identity(self):
        records = fixture()
        for key in list(records):
            if key.startswith(f"ManagedElement={LTE},"):
                records[key.replace("ExternalGNodeBFunction=" + NR, "ExternalGNodeBFunction=5152-000902007")
                        .replace("TermPointToGNB=" + NR, "TermPointToGNB=5152-000902007")] = records.pop(key)
        row, = audit_termpoint_to_gnb(records, nodes=[LTE])
        self.assertEqual(row.status, "Mismatch")

    def test_other_sites_and_neighbor_terms_ignored(self):
        records = fixture()
        records[f"ManagedElement={LTE},ENodeBFunction=1,GUtraNetwork=1,ExternalGNodeBFunction=OTHER,TermPointToGNB=OTHER"] = {"ipAddress": "0.0.0.0"}
        records[f"ManagedElement=MIN31170_SITEB02,{TERM}"] = {"ipAddress": "0.0.0.0"}
        self.assertEqual(len(audit_termpoint_to_gnb(records)), 2)

    def test_no_nr_no_audit(self):
        records = {k: v for k, v in fixture().items() if "GNBDUFunction=" not in k}
        self.assertEqual(audit_termpoint_to_gnb(records), [])

    def test_missing_expected_ip_not_match_or_correction(self):
        records = fixture()
        del records[f"ManagedElement={NR},{ADDRESS}"]
        self.assertEqual({r.status for r in audit_termpoint_to_gnb(records)}, {"NotFound"})

    def test_missing_termpoint(self):
        records = fixture()
        del records[f"ManagedElement={LTE},{TERM}"]
        row, = audit_termpoint_to_gnb(records, nodes=[LTE])
        self.assertEqual(row.status, "MO_NotFound")

    def test_two_nr_targets_keep_their_own_ip(self):
        records = fixture()
        other = "MIN3117_SITEB03"
        records[f"ManagedElement={other},GNBDUFunction=1"] = {"gNBId": "99"}
        records[f"ManagedElement={other},{ADDRESS}"] = {"address": "10.40.128.2/31"}
        records[f"ManagedElement={LTE},{TERM.replace(NR, other)}"] = {"ipAddress": "10.40.128.2"}
        rows = audit_termpoint_to_gnb(records, nodes=[LTE])
        self.assertEqual([(r.ref_cell, r.expected, r.status) for r in rows],
                         [(NR, "10.40.128.0", "Mismatch"), (other, "10.40.128.2", "Match")])

    def test_audit_only_no_generated_set_commands(self):
        from audit.audit_core import _NON_SETTABLE_CATEGORIES
        self.assertIn("termpoint-gnb", _NON_SETTABLE_CATEGORIES)

    def test_intmom_modump(self):
        payload = "\n".join(f"{dn} {k} {v}" for dn, attrs in fixture().items()
                            for k, v in attrs.items())
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "intmomlog.txt"
            path.write_text(payload)
            rows = audit_termpoint_to_gnb(parse_dump(str(path)))
        self.assertEqual([r.status for r in rows], ["Match", "Mismatch"])

    def test_excel_contains_target_actual_and_expected(self):
        from audit.audit_core import write_excel
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "audit.xlsx"
            write_excel(audit_termpoint_to_gnb(fixture()), str(path), {})
            book = load_workbook(path)
            rows = [row for sheet in book for row in sheet.iter_rows(values_only=True)
                    if "termpoint-gnb" in row and "Mismatch" in row]
            self.assertEqual(len(rows), 1)
            for value in (NR, LTE, "10.40.128.0", "0.0.0.0"):
                self.assertIn(value, rows[0])
            book.close()


if __name__ == "__main__":
    unittest.main()
