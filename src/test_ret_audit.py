"""RET audit: .ims2 reader, Ericsson lhgetc parser, before/after grading."""
import gzip
import io
import os
import struct
import tempfile
import unittest
import zipfile

from audit import ret_audit as ra
import ims2_reader

HERE = os.path.dirname(os.path.abspath(__file__))
LOG = os.path.join(HERE, "testdata", "ret_audit",
                   "MIN1148_lhgetc_antennanearunit.log")

META_XML = b"""<?xml version="1.0" encoding="UTF-8"?>
<infoModel id="lte">
  <managedObject class="ALD_R">
    <p name="operationalState" type="EOp"><proto index="1" repeated="false" type="string"/></p>
    <p name="serialNumber" type="string"><proto index="2" repeated="false" type="string"/></p>
    <p name="vendorCode" type="string"><proto index="3" repeated="false" type="string"/></p>
    <p name="productCode" type="string"><proto index="4" repeated="false" type="string"/></p>
  </managedObject>
  <managedObject class="RETU_R">
    <enumeration name="EOp"><enum name="disabled" value="0"/><enum name="enabled" value="1"/></enumeration>
    <struct name="AntBand">
      <p name="antFreqBand" type="integer"><proto index="1" repeated="false" type="int32"/></p>
      <p name="antOperGain" type="integer"><proto index="2" repeated="false" type="int32"/></p>
    </struct>
    <p name="operationalState" type="EOp"><proto index="1" repeated="false" type="enum"/></p>
    <p name="angle" type="integer"><proto index="2" repeated="false" type="int32"/></p>
    <p name="subunitNumber" type="integer"><proto index="3" repeated="false" type="int32"/></p>
    <p name="baseStationID" type="string"><proto index="4" repeated="false" type="string"/></p>
    <p name="configDN" type="string"><proto index="5" repeated="false" type="string"/></p>
    <p name="antBandList" type="AntBand" recurrence="repeated"><proto index="6" repeated="true" type="Length-Delimited"/></p>
    <p name="minAngle" type="integer"><proto index="7" repeated="false" type="sint32"/></p>
    <p name="sectorID" type="string"><proto index="8" repeated="false" type="string"/></p>
  </managedObject>
  <managedObject class="ALARM">
    <p name="alarmNumber" type="integer"><proto index="1" repeated="false" type="int32"/></p>
    <p name="alarmName" type="string"><proto index="2" repeated="false" type="string"/></p>
    <p name="alarmSeverity" type="string"><proto index="3" repeated="false" type="string"/></p>
    <p name="observationTime" type="string"><proto index="4" repeated="false" type="string"/></p>
    <p name="alarmingResourceDN" type="string"><proto index="5" repeated="false" type="string"/></p>
  </managedObject>
</infoModel>
"""
MAGIC = ims2_reader.MAGIC


def _varint(n):
    out = b""
    while True:
        b = n & 0x7F
        n >>= 7
        out += bytes([b | (0x80 if n else 0)])
        if not n:
            return out


def _f(idx, wire, payload):
    tag = _varint(idx << 3 | wire)
    if wire == 0:
        return tag + _varint(payload)
    return tag + _varint(len(payload)) + payload


def _record(dn, payload, flag=0):
    dn = dn.encode()
    if flag == 1:                    # delete: no length, no payload
        return struct.pack(">H", len(dn)) + dn + bytes([1])
    return struct.pack(">H", len(dn)) + dn + bytes([0]) + \
        struct.pack(">I", len(payload)) + payload


DELETE = object()      # a record that deletes the MO (flag 1, no payload)


def make_ims2(path, records_per_member):
    zbuf = io.BytesIO()
    with zipfile.ZipFile(zbuf, "w") as z:
        z.writestr("lte/meta.xml", META_XML)
    zbody = zbuf.getvalue()
    out = struct.pack(">III", 1, 22, 18) + b"BTSOAM IM SNAPSHOT" + MAGIC
    out += struct.pack(">IQ", 2, len(zbody)) + zbody + MAGIC
    out += struct.pack(">IQ", 3, 4) + b"\x00\x00\x00\x01" + MAGIC
    data = b""
    for i, records in enumerate(records_per_member):
        block = struct.pack(">III", 413, 1000 + i, len(records))
        block += b"".join(_record(dn, b"", 1) if p is DELETE
                          else _record(dn, p) for dn, p in records)
        member = gzip.compress(block)
        if i:
            data += MAGIC + struct.pack(">IQ", 0, len(member))
        data += member
    out += struct.pack(">IQ", 0, 0xFFFFFFFF) + data
    with open(path, "wb") as fh:
        fh.write(out)


RV = "/MRBTS-1/RAT-1/RUNTIME_VIEW-1/MRBTS_R-1/EQM_R-1/APEQM_R-1"


class Ims2ReaderTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.path = os.path.join(self.tmp, "im-snapshot_MRBTS-342253_X.ims2")
        ald = (_f(1, 2, b"enabled") + _f(2, 2, b"K1-M1") + _f(3, 2, b"RB")
               + _f(4, 2, b"SCU-002AL"))
        band = _f(1, 0, 8) + _f(2, 0, 150)
        retu_old = _f(1, 0, 1) + _f(2, 0, 30) + _f(3, 0, 1)
        retu = (_f(1, 0, 1) + _f(2, 0, 20) + _f(3, 0, 1)
                + _f(4, 2, b"342253_AY1") + _f(5, 2, b"/NP-1/ALD_A-1/RETU_A-1")
                + _f(6, 2, band) + _f(6, 2, _f(1, 0, 12) + _f(2, 0, 146))
                + _f(7, 0, 39)          # sint32 zigzag -20
                + _f(8, 2, b"SEC-2_ANT1&2")
                + _f(99, 0, 7))         # unknown field kept as #99
        FRI = "/MRBTS-1/RAT-1/FRI-1"

        def alarm(no, name, t, obj):
            return (_f(1, 0, no) + _f(2, 2, name) + _f(3, 2, b"Major")
                    + _f(4, 2, t) + _f(5, 2, obj))
        ret_obj = RV.encode() + b"/ALD_R-1/RETU_R-1"
        make_ims2(self.path, [
            [(RV + "/ALD_R-1", ald), (RV + "/ALD_R-1/RETU_R-1", retu_old),
             (FRI + "/ALARM-1", alarm(7103, b"FAN", b"20260927100000.0+0800",
                                      b"/x/EAC_R-7")),
             (FRI + "/ALARM-9", alarm(7100, b"GONE", b"20260101000000.0+0800",
                                      b"/x/Y"))],
            [(RV + "/ALD_R-1/RETU_R-1", retu),      # later write wins
             (FRI + "/ALARM-1", DELETE),            # FAN alarm cleared
             (FRI + "/ALARM-2", alarm(7113, b"RET FAIL",
                                      b"20260928120000.0+0800", ret_obj)),
             (FRI + "/ALARM-9", DELETE)],
        ])

    def test_decodes_with_model_and_last_write_wins(self):
        snap = ims2_reader.Snapshot(self.path)
        v = snap.get(RV + "/ALD_R-1/RETU_R-1")
        self.assertEqual(v["operationalState"], "enabled")
        self.assertEqual(v["angle"], 20)
        self.assertEqual(v["minAngle"], -20)
        self.assertEqual(v["baseStationID"], "342253_AY1")
        self.assertEqual([b["antOperGain"] for b in v["antBandList"]],
                         [150, 146])
        self.assertEqual(v["#99"], 7)

    def test_nokia_rets(self):
        rets = ra.nokia_rets(ims2_reader.Snapshot(self.path))
        self.assertEqual(len(rets), 1)
        r = rets[0]
        self.assertEqual(r.node, "MRBTS-342253")
        self.assertEqual((r.serial, r.subunit, r.tilt), ("K1-M1", 1, 20))
        self.assertEqual(r.unique_id, "RBK1-M1")
        self.assertEqual(r.gains, "150 146")
        # no radio link in this snapshot: sector falls back to the RET's own
        # Sector ID, so it cannot be cross-checked against itself
        self.assertEqual((r.sector, r.sector_check), ("S2", ""))
        # the active RET alarm makes it faulty before the swap
        self.assertEqual(r.faults, ["alarm 7113 RET FAIL"])

    def test_deleted_mo_is_gone(self):
        snap = ims2_reader.Snapshot(self.path)
        self.assertNotIn("/MRBTS-1/RAT-1/FRI-1/ALARM-1", snap.raw)
        self.assertEqual(
            [(dn.rsplit("/", 1)[1], deleted) for dn, deleted, _p in snap.history],
            [("ALARM-1", False), ("ALARM-9", False), ("ALARM-1", True),
             ("ALARM-2", False), ("ALARM-9", True)])

    def test_alarms_active_and_history(self):
        active, history = ra.nokia_alarms(ims2_reader.Snapshot(self.path))
        self.assertEqual([a["Alarm"] for a in active], ["7113"])
        self.assertEqual(active[0]["Raised"], "2026-09-28 12:00:00 +0800")
        self.assertEqual([(a["Alarm"], a["Status"]) for a in history],
                         [("7113", "Active"), ("7103", "Cleared"),
                          ("7100", "Cleared")])

    def test_rejects_other_files(self):
        bad = os.path.join(self.tmp, "x.ims2")
        with open(bad, "wb") as fh:
            fh.write(b"PK\x03\x04 not a snapshot")
        with self.assertRaises(ims2_reader.Ims2Error):
            ims2_reader.Snapshot(bad)


class SiteCheckTests(unittest.TestCase):
    def test_common_core_of_real_names(self):
        self.assertEqual(ra.common_core("TCFGACROSRUBTUPISCOTFWLY",
                                        "MIN1148_CROSRUBTUPISCOTB01"),
                         "CROSRUBTUPISCOT")

    def test_same_site(self):
        c = ra.site_check("TCFGACROSRUBTUPISCOTFWLY",
                          ["MIN1148_CROSRUBTUPISCOTB01",
                           "MIN1148_CROSRUBTUPISCOTB02"])
        self.assertTrue(c["ok"])
        self.assertEqual([core for _n, core in c["nodes"]],
                         ["CROSRUBTUPISCOT", "CROSRUBTUPISCOT"])

    def test_one_node_from_another_site_fails(self):
        c = ra.site_check("TCFGACROSRUBTUPISCOTFWLY",
                          ["MIN1148_CROSRUBTUPISCOTB01",
                           "MIN2207_SANJUANPOBLACIONB01"])
        self.assertFalse(c["ok"])
        self.assertIn("MIN2207_SANJUANPOBLACIONB01", c["message"])

    def test_missing_names_fail(self):
        self.assertFalse(ra.site_check("", ["MIN1148_X"])["ok"])
        self.assertFalse(ra.site_check("TCFGACROSRUBTUPISCOTFWLY", [])["ok"])

    def test_bts_name_falls_back_to_file_name(self):
        class Snap:
            path = ("im-snapshot_MRBTS-342253_TCFGACROSRUBTUPISCOTFWLY_"
                    "SBTS25R2_ENB_0000.ims2")

            def by_class(self, *cls):
                return iter(())
        self.assertEqual(ra.nokia_bts_name(Snap()), "TCFGACROSRUBTUPISCOTFWLY")


class EricssonParserTests(unittest.TestCase):
    def setUp(self):
        with open(LOG, encoding="utf-8") as fh:
            self.rets = ra.ericsson_rets(fh.read())
        self.by = {(r.serial, r.subunit): r for r in self.rets}

    def test_all_subunits_with_node(self):
        self.assertEqual(len(self.rets), 24)
        self.assertEqual({r.node for r in self.rets},
                         {"MIN1148_CROSRUBTUPISCOTB01"})

    def test_values_and_continuation_rows(self):
        r = self.by[("K77222C4200104-M1", 1)]      # the verbatim padded row
        self.assertEqual(r.unique_id, "RBK77222C4200104-M1")
        self.assertEqual(r.device_id, "RBK77222C4200104-M1")
        self.assertEqual((r.product, r.hw, r.sw),
                         ("SCU-002AL", "2.00", "MAL1.01"))
        self.assertEqual((r.tilt, r.min_tilt, r.max_tilt), (120, 20, 120))
        self.assertEqual(r.gains, "158 150 164 160")
        self.assertEqual(r.base_station_id, "342253_AY1L1800&L2100&G1800")
        self.assertEqual(r.state, "DISABLED")
        self.assertIn("AntennaNearUnit DISABLED", r.faults)
        self.assertIn("availability FAILED", r.faults)
        ok = self.by[("K77222C4200218-M2", 3)]
        self.assertEqual(ok.unique_id, "M2")
        self.assertEqual((ok.state, ok.faults), ("ENABLED", []))

    def test_value_normalisation(self):
        self.assertEqual(ra._value("1 (UNLOCKED)  "), "UNLOCKED")
        self.assertEqual(ra._value("i[1] = 1 (FAILED) "), "FAILED")
        self.assertEqual(ra._value("i[0] =            "), "")
        self.assertEqual(ra._value("i[4] = 158 150"), "158 150")
        self.assertEqual(ra._value("[1] = AntennaUnitGroup=X"),
                         "AntennaUnitGroup=X")


# An AntennaNearUnit configured on Ericsson whose device never answered: no
# serial / onUnitUniqueId, empty iuant* data, only the configured uniqueId.
UNDETECTED_LOG = "\n".join([
    "NODEX> lhgetc AntennaUnitGroup=.*,AntennaNearUnit=",
    "MO;administrativeState;operationalState;serialNumber;onUnitUniqueId;"
    "uniqueId;rfPortRef",
    "AntennaUnitGroup=Site_S2,AntennaNearUnit=1;1 (UNLOCKED);0 (DISABLED);;;"
    "RBK9-M1;FieldReplaceableUnit=RRU2,RfPort=R",
    "MO;availabilityStatus;electricalAntennaTilt;iuantBaseStationId;"
    "operationalState;retSubUnitId;subunitNumber",
    "AntennaUnitGroup=Site_S2,AntennaNearUnit=1,RetSubUnit=1;"
    "i[1] = 1 (FAILED);;;0 (DISABLED);1;",
])


def _nokia_like(e, **over):
    r = ra.Ret(**{k: getattr(e, k) for k in ra.Ret.__dataclass_fields__
                  if k != "faults"})
    r.side, r.faults = "Nokia", []
    for k, v in over.items():
        setattr(r, k, v)
    return r


class CompareTests(unittest.TestCase):
    def setUp(self):
        with open(LOG, encoding="utf-8") as fh:
            self.eric = {(r.serial, r.subunit): r
                         for r in ra.ericsson_rets(fh.read())}
        self.ok = self.eric[("K77222C4200218-M1", 1)]      # ENABLED
        self.bad = self.eric[("K77222C4200110-M1", 1)]     # DISABLED

    # ── target: ENABLED on Ericsson, same sector ────────────────────────
    def test_enabled_same_sector_is_ok(self):
        rows = ra.compare([_nokia_like(self.ok)], [self.ok])
        self.assertEqual((rows[0].status, rows[0].diffs), ("OK", []))
        self.assertFalse(rows[0].is_issue)

    def test_enabled_with_differences_is_still_ok(self):
        n = _nokia_like(self.ok, tilt=20, ant_model="OTHER")
        row = ra.compare([n], [self.ok])[0]
        self.assertEqual(row.status, "OK")
        self.assertEqual(row.diffs, ["Tilt (deg)", "Antenna model"])
        self.assertIn("mismatch Antenna model", row.remark)
        self.assertIn("tilt 2.0° → 12.0° (note)", row.remark)
        self.assertNotIn("mismatch Tilt", row.remark)

    def test_enabled_but_sector_changed_is_not_ok(self):
        n = _nokia_like(self.ok, sector="S1")
        row = ra.compare([n], [self.ok])[0]
        self.assertEqual(row.status, "Not OK")
        self.assertIn("sector changed S1 → S3", row.remark)

    def test_disabled_on_ericsson_is_not_ok(self):
        n = _nokia_like(self.bad, state="ENABLED")
        row = ra.compare([n], [self.bad])[0]
        self.assertEqual(row.status, "Not OK")
        self.assertTrue(row.remark.startswith(
            "Ericsson DISABLED (AntennaNearUnit DISABLED"))

    def test_disabled_on_both_is_still_not_ok(self):
        n = _nokia_like(self.bad, faults=["RETU disabled"])
        row = ra.compare([n], [self.bad])[0]
        self.assertEqual(row.status, "Not OK")
        self.assertIn("Nokia already: RETU disabled", row.remark)

    def test_faulty_on_nokia_enabled_on_ericsson_is_ok(self):
        n = _nokia_like(self.ok, state="DISABLED", faults=["RETU disabled"])
        self.assertEqual(ra.compare([n], [self.ok])[0].status, "OK")

    # ── unique ID: uniqueId first, then onUnitUniqueId ──────────────────
    def test_unique_id_falls_back_to_on_unit_unique_id(self):
        e = self.eric[("K77222C4200218-M2", 3)]    # uniqueId 'M2' only
        self.assertEqual(e.unique_id, "M2")
        n = _nokia_like(e, unique_id="RBK77222C4200218-M2")
        self.assertTrue(ra.unique_id_matches(n, e))
        self.assertNotIn("Unique ID", ra.compare([n], [e])[0].diffs)
        # the report shows the value that matched: onUnitUniqueId
        self.assertEqual(ra.ericsson_unique_id(n, e), "RBK77222C4200218-M2")
        # uniqueId itself is shown when it already matches
        full = self.eric[("K77222C4200104-M1", 1)]
        self.assertEqual(
            ra.ericsson_unique_id(_nokia_like(full), full),
            full.unique_id)
        # an Ericsson-only RET shows its uniqueId
        self.assertEqual(ra.ericsson_unique_id(None, e), "M2")

    def test_unique_id_different_on_both_is_red(self):
        n = _nokia_like(self.ok, unique_id="RBOTHER-M1")
        row = ra.compare([n], [self.ok])[0]
        self.assertIn("Unique ID", row.diffs)
        self.assertIn("mismatch Unique ID", row.remark)
        self.assertEqual(row.status, "OK")        # enabled still wins

    # ── one side only ───────────────────────────────────────────────────
    def test_nokia_only_is_not_ok(self):
        row = ra.compare([_nokia_like(self.ok)], [])[0]
        self.assertEqual(row.status, "Not OK")
        self.assertIn("not found on Ericsson", row.remark)

    def test_ericsson_only_enabled_is_ok(self):
        row = ra.compare([], [self.ok])[0]
        self.assertEqual(row.status, "OK")
        self.assertIn("new RET", row.remark)

    def test_ericsson_only_disabled_is_not_ok(self):
        self.assertEqual(ra.compare([], [self.bad])[0].status, "Not OK")

    def test_ericsson_mo_without_device_pairs_on_unique_id(self):
        rets = ra.ericsson_rets(UNDETECTED_LOG)
        self.assertEqual(len(rets), 1)
        e = rets[0]
        self.assertEqual((e.serial, e.sector), ("", "S2"))
        self.assertIn("device not detected", e.faults)
        n = _nokia_like(self.ok, serial="K9-M1", device_id="RBK9-M1",
                        unique_id="RBK9-M1", subunit=1, sector="S2")
        row = ra.compare([n], rets)[0]
        self.assertEqual((row.status, row.paired_by),
                         ("Not OK", "configured uniqueId"))

    # ── replaced antenna (different device, same position) ──────────────
    def _replaced_pair(self, **eric_over):
        n = _nokia_like(self.ok, serial="OLD-M1", device_id="RBOLD-M1",
                        unique_id="RBOLD-M1", ant_serial="OLDANT",
                        sector="S3")
        e = _nokia_like(self.ok, **eric_over)
        e.side = "Ericsson"
        e.faults = list(eric_over.get("faults", []))
        return n, e

    def test_replaced_enabled_is_ok(self):
        n, e = self._replaced_pair()
        row = ra.compare([n], [e])[0]
        self.assertEqual((row.status, row.paired_by), ("OK", "base station ID"))
        self.assertIn("different device, paired by base station ID", row.remark)
        self.assertIn("mismatch Unique ID, Antenna serial", row.remark)

    def test_replaced_disabled_is_not_ok(self):
        n, e = self._replaced_pair(state="DISABLED",
                                   faults=["RetSubUnit DISABLED"])
        self.assertEqual(ra.compare([n], [e])[0].status, "Not OK")

    def test_replaced_pairs_on_sector_id_when_bsid_changed(self):
        n, e = self._replaced_pair(base_station_id="NEWBSID")
        row = ra.compare([n], [e])[0]
        self.assertEqual((row.status, row.paired_by), ("OK", "sector ID"))
        self.assertIn("Base station ID", row.diffs)

    def test_sector_id_cross_check(self):
        r = _nokia_like(self.ok, sector="S1", sector_src="RMOD_R-4 cells",
                        sector_id="SEC-2_ANT1&2")
        self.assertEqual(ra.sector_check(r),
                         "Sector ID SEC-2_ANT1&2 but RET is on S1")
        r.sector = "S2"
        self.assertEqual(ra.sector_check(r), "OK")
        self.assertEqual(self.ok.sector_check, "OK")   # SEC-3 on S3
        # a wrong programmed Sector ID is a note only
        n = _nokia_like(self.ok, sector_id="SEC-9_X")
        n.sector_check = ra.sector_check(n)
        row = ra.compare([n], [self.ok])[0]
        self.assertEqual(row.status, "OK")
        self.assertIn("Nokia: Sector ID SEC-9_X but RET is on S3", row.remark)

    def test_ericsson_sector_from_group_name(self):
        self.assertEqual(self.ok.sector, "S3")
        self.assertEqual(
            ra.ericsson_sector("AntennaUnitGroup=X,AntennaNearUnit=1", {
                "rfPortRef": "FieldReplaceableUnit=B1B3_RRU2,RfPort=R"}, {})[0],
            "S2")

    def test_excel_is_written(self):
        rows = ra.compare([_nokia_like(self.ok, tilt=20)], [self.ok, self.bad])
        path = os.path.join(tempfile.mkdtemp(), "ret.xlsx")
        alarms = ([{"Status": "Active", "Alarm": "7113", "Severity": "Major"}],
                  [{"Status": "Cleared", "Alarm": "7103", "Severity": "Minor"}])
        ra.write_excel(rows, [], [self.ok, self.bad], path,
                       ra.default_meta("a.ims2", ["b.log"], "SITE"),
                       alarms=alarms)
        from openpyxl import load_workbook
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames,
                         ["Summary", "RET Audit", "Nokia RET", "Ericsson RET",
                          "Nokia Alarms"])
        al = wb["Nokia Alarms"]
        self.assertEqual(al.cell(1, 1).value, "ACTIVE ALARMS (1)")
        self.assertEqual(al.cell(3, 2).value, "7113")
        self.assertEqual(al.cell(5, 1).value,
                         "ALARM HISTORY (from the snapshot log, newest first) (1)")
        self.assertEqual((al.cell(7, 1).value, al.cell(7, 2).value),
                         ("Cleared", "7103"))
        ws = wb["RET Audit"]
        self.assertEqual(ws.cell(1, 1).value, "Status")
        self.assertEqual(ws.max_row, 3)


if __name__ == "__main__":
    unittest.main()
