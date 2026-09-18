import tempfile
import unittest
from pathlib import Path
from audit.dump_parser import parse_dump
from audit.power_audit import audit_power_license
from test_audit_termpoint import cmdump


def fixture(groups, node='BB1', grant='61', serial=False):
    records = {f'ManagedElement={node},Lm=1,CapacityState=CXC4012338': {'grantedCapacityLevel': grant}}
    index = 0
    for radio, powers in enumerate(groups):
        fru = f'Equipment=1,FieldReplaceableUnit=R{radio}'
        records[f'ManagedElement={node},{fru}'] = {'serialNumber': f'SERIAL{radio}'} if serial else {}
        sef = f'SectorEquipmentFunction=R{radio}'
        branch = f'RfBranch=R{radio}'
        records[f'ManagedElement={node},{sef}'] = {'rfBranchRef': branch}
        records[f'ManagedElement={node},{branch}'] = {'rfPortRef': fru + ',RfPort=A'}
        for power in powers:
            cls = ('SectorCarrier', 'NRSectorCarrier', 'Trx')[index % 3]
            records[f'ManagedElement={node},{cls}={index}'] = {
                'configuredMaxTxPower': str(power),
                'sectorFunctionRef' if cls == 'SectorCarrier' else 'sectorEquipmentFunctionRef': sef}
            index += 1
    return records


class PowerAuditTests(unittest.TestCase):
    def test_user_example_both_formats(self):
        records = fixture([[160000,160000,80000,80000,80000], [160000,80000,40000],
                           [80000,80000], [80000,80000], [80000,80000]])
        dcg = '\n'.join('MO  '+dn+'\n'+'\n'.join(f'{k}  {v}' for k,v in attrs.items())
                        for dn,attrs in records.items()).encode()
        with tempfile.TemporaryDirectory() as folder:
            for kind,payload in (('xml',cmdump(records)),('log',dcg)):
                with self.subTest(kind=kind):
                    path = Path(folder)/('dump.'+kind)
                    path.write_bytes(payload)
                    rows,evidence = audit_power_license(parse_dump(str(path)))
                    self.assertEqual((rows[0].expected,rows[0].actual,rows[0].status),('61','61','Match'))
                    self.assertEqual(len(evidence),14)
                    self.assertIn('560 W -> 27 units',rows[0].remark)
                    self.assertIn('280 W -> 13 units',rows[0].remark)
                    self.assertEqual(rows[0].remark.count('160 W -> 7 units'),3)
        records['ManagedElement=BB1,Lm=1,CapacityState=CXC4012338']['grantedCapacityLevel']='60'
        rows,_=audit_power_license(records)
        self.assertEqual(rows[0].status,'Mismatch')
        self.assertIn('shortage 1',rows[0].remark)

    def test_shared_radio_discount_per_bb(self):
        records=fixture([[40000,40000]],grant='3',serial=True)
        records.update(fixture([[40000,40000]],node='BB2',grant='3',serial=True))
        rows,_=audit_power_license(records)
        self.assertEqual([(r.expected,r.status) for r in rows],[('3','Match'),('3','Match')])

    def test_minimum_and_exact_precision(self):
        for power,expected in ((0,'0'),(10000,'0'),(20000,'0'),(20001,'0.00005'),(426000,'20.3')):
            rows,_=audit_power_license(fixture([[power]],grant='0'))
            self.assertEqual(rows[0].expected,expected)
            self.assertEqual(rows[0].status,'Match' if power<=20000 else 'Mismatch')

    def test_alias_and_duplicate_branches_discount_once(self):
        records=fixture([[40000],[40000]],grant='3',serial=True)
        records['ManagedElement=BB1,Equipment=1,FieldReplaceableUnit=R1']['serialNumber']='SERIAL0'
        records['ManagedElement=BB1,SectorEquipmentFunction=R0']['rfBranchRef']='RfBranch=R0;RfBranch=R0'
        rows,_=audit_power_license(records)
        self.assertEqual((rows[0].expected,rows[0].status),('3','Match'))

    def test_incomplete_data_never_passes(self):
        for value in ('','-1','NaN','Infinity','not-a-number'):
            rows,_=audit_power_license(fixture([[value]],grant='10'))
            self.assertEqual(rows[0].status,'NotFound')
        for missing in ('SectorEquipmentFunction=R0','RfBranch=R0','Equipment=1,FieldReplaceableUnit=R0'):
            records=fixture([[80000]])
            del records['ManagedElement=BB1,'+missing]
            rows,_=audit_power_license(records)
            self.assertEqual(rows[0].status,'NotFound')
        rows,_=audit_power_license({},nodes=['MISSING'])
        self.assertEqual(rows[0].status,'NotFound')

    def test_multiple_radios_or_wrong_owner_unresolved(self):
        records=fixture([[80000],[80000]])
        records['ManagedElement=BB1,SectorEquipmentFunction=R0']['rfBranchRef']='RfBranch=R0;RfBranch=R1'
        rows,_=audit_power_license(records)
        self.assertEqual(rows[0].status,'NotFound')
        self.assertIn('multiple physical radios',rows[0].remark)
        records=fixture([[80000]])
        records['ManagedElement=BB1,SectorCarrier=0']['sectorFunctionRef']='ManagedElement=OTHER,SectorEquipmentFunction=R0'
        rows,_=audit_power_license(records)
        self.assertEqual(rows[0].status,'NotFound')

    def test_excel_and_no_license_correction(self):
        from audit.audit_core import write_excel,generate_moshell_scripts,generate_cmedit_scripts
        from openpyxl import load_workbook
        rows,evidence=audit_power_license(fixture([[426000]],grant='20'))
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'report.xlsx'
            write_excel(rows,str(path),{},power_results=rows,power_evidence=evidence)
            book=load_workbook(path)
            self.assertNotIn('Power License',book.sheetnames)
            self.assertNotIn('Power MO Detail',book.sheetnames)
            self.assertEqual(book['Detail']['F2'].value,'20.3')
            self.assertEqual(book['Detail']['G2'].value,'20')
            self.assertIn('shortage 0.3',book['Detail']['J2'].value)
            self.assertIn('426 W -> 20.3 units',book['Detail']['J2'].value)
            book.close()
            self.assertEqual(generate_moshell_scripts(rows,folder,'SITE',str(path)),[])
            self.assertEqual(generate_cmedit_scripts(rows,folder,'SITE',str(path)),[])

if __name__=='__main__':
    unittest.main()
