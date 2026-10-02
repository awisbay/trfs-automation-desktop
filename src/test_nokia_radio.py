"""Nokia radio view: VSWR per radio/port/band + RTWP/RSSI per cell/TRX."""
import os
import tempfile
import unittest

from audit import nokia_radio as nr
from audit import ret_audit as ra

RV = "/MRBTS-1/RAT-1/RUNTIME_VIEW-1/MRBTS_R-1/EQM_R-1/APEQM_R-1"
EQ = "/MRBTS-1/RAT-1/BTS_L-1/EQM_L-1"
NP = "/MRBTS-1/RAT-1/BTS_L-1/BTS_CONF-1/NP-7/SCF-1/MRBTS_A-9"
MAP = NP + "/MNL_A-1/MNLENT_A-1/CELLMAPPING_A-1"
MC = "/MRBTS-1/RAT-1/MCTRL-1/BBTOP_M-1/MRBTS_M-1"


class FakeSnap:
    def __init__(self, data):
        self.data = data

    def by_class(self, *classes):
        for cls in classes:
            for dn, v in self.data.get(cls, []):
                yield dn, v


def snapshot():
    ant = NP + "/EQM_A-1/APEQM_A-1/RMOD_A-2/ANTL_A-%d"
    return FakeSnap({
        # runtime radio 2 (plan RMOD_A-2) is local RMOD_L-5: matched on serial
        "RMOD_R": [(RV + "/RMOD_R-2", {"serialNumber": "EA1", "productName":
                                       "AHPDA", "configDN": NP +
                                       "/EQM_A-1/APEQM_A-1/RMOD_A-2"})],
        "RMOD_L": [(EQ + "/RMOD_L-5", {"serialNumber": "EA1"})],
        # filters 1-2 = B8 (pipe a), 3-4 = B28 (pipe b)
        "FFU_L": [(EQ + "/RMOD_L-5/RU_L-1/FF_L-%d/FFU_L-1" % i,
                   {"centerFrequencyDownlink": f})
                  for i, f in ((3, 780.5e6), (1, 947.5e6), (2, 947.5e6),
                               (4, 780.5e6))],
        "VSWR": [(EQ + "/RMOD_L-5/RU_L-1/VUBUS_L-%s/VU_L-1/VSWR-%s" % (p, k),
                  {"rp1Name": "antenna%s%s" % (p, pipe), "vswr": val,
                   "vswrMinorLimit": 15.0, "vswrMajorLimit": 17.0})
                 for p, k, pipe, val in ((1, 1, "a", 13.0), (1, 2, "b", 12.0),
                                         (2, 1, "a", 18.0), (2, 2, "b", 15.0),
                                         (3, 1, "a", 14.0))],
        "LNCEL_A": [(NP + "/LNBTS_A-9/LNCEL_A-122", {"cellName": "SITE-Y-122"})],
        "LNCEL_FDD_A": [(NP + "/LNBTS_A-9/LNCEL_A-122/LNCEL_FDD_A-1",
                         {"earfcnDL": 3749})],          # B8
        "LCELC_A": [(MAP + "/LCELC_A-2", {"bandNumber": 8})],
        "CHANNEL_A": [
            (MAP + "/LCELL_A-122/CHANNELGROUP_A-1/CHANNEL_A-2",
             {"antlDN": ant % 1, "direction": "RX"}),
            (MAP + "/LCELL_A-122/CHANNELGROUP_A-1/CHANNEL_A-3",
             {"antlDN": ant % 2, "direction": "RX"}),
            (MAP + "/LCELL_A-122/CHANNELGROUP_A-1/CHANNEL_A-1",
             {"antlDN": ant % 1, "direction": "TX"}),
            (MAP + "/LCELC_A-2/CHANNELGROUP_A-3/CHANNEL_A-2",
             {"antlDN": ant % 1, "direction": "RX"}),
        ],
        "RTWP_MEASUREMENT": [
            (MC + "/LNBTS_M-1/CELL_M-122/CHANNEL_GROUP_M-1/CHANNEL_M-2/"
                  "RTWP_MEASUREMENT-1", {"rtwpValue": -950.0}),
            (MC + "/LNBTS_M-1/CELL_M-122/CHANNEL_GROUP_M-1/CHANNEL_M-3/"
                  "RTWP_MEASUREMENT-1", {"rtwpValue": -990.0}),
        ],
        "RSSI_MEASUREMENT": [
            (MC + "/GNBTS_M-1/CELL_M-2003/CHANNEL_GROUP_M-3/CHANNEL_M-2/"
                  "RSSI_MEASUREMENT-1", {"rssiValue": -71}),
        ],
    })


class RadioViewTests(unittest.TestCase):
    def setUp(self):
        self.rows = nr.radio_rows(snapshot())

    def row(self, port, band, cell=""):
        return next(r for r in self.rows if r["Port"] == port
                    and r["Band"] == band and str(r["Cell"]).startswith(cell))

    def test_local_radio_matched_on_serial(self):
        self.assertEqual({(r["Radio"], r["Product"], r["Sector"])
                          for r in self.rows}, {("RMOD-2", "AHPDA", "S2")})

    def test_pipe_a_is_first_filter_band(self):
        self.assertEqual(self.row("ANT1", "B8", "SITE")["VSWR"], 1.3)
        self.assertEqual(self.row("ANT1", "B28")["VSWR"], 1.2)

    def test_vswr_status_uses_radio_limits(self):
        self.assertEqual(self.row("ANT2", "B8", "SITE")["VSWR status"], "Major")
        self.assertEqual(self.row("ANT2", "B28")["VSWR status"], "Minor")
        self.assertEqual(self.row("ANT1", "B8", "SITE")["VSWR status"], "OK")
        self.assertEqual(self.row("ANT3", "B8")["VSWR status"], "Warning")

    def test_lte_rtwp_per_branch_on_its_band(self):
        r1 = self.row("ANT1", "B8", "SITE")
        self.assertEqual((r1["Measure"], r1["RX (dBm)"]), ("RTWP", -95.0))
        self.assertEqual(self.row("ANT2", "B8", "SITE")["RX (dBm)"], -99.0)
        self.assertEqual(r1["Branch imbalance (dB)"], 4.0)

    def test_gsm_rssi_per_trx(self):
        g = self.row("ANT1", "B8", "GSM")
        self.assertEqual((g["TRX"], g["Measure"], g["RX (dBm)"]),
                         (3, "RSSI", -71.0))

    def test_port_band_without_cell_keeps_vswr_row(self):
        b28 = self.row("ANT2", "B28")
        self.assertEqual((b28["Cell"], b28["VSWR"]), ("", 1.5))

    def test_sheet_written(self):
        path = os.path.join(tempfile.mkdtemp(), "r.xlsx")
        ra.write_excel([], [], [], path, {"Site": "X"}, radio=self.rows)
        from openpyxl import load_workbook
        ws = load_workbook(path)["Nokia Radio"]
        self.assertEqual(ws.cell(1, 1).value, "Sector")
        self.assertEqual(ws.max_row, len(self.rows) + 3)   # + note row


RV3 = "/MRBTS-1/RAT-1/RUNTIME_VIEW-3/MRBTS_R-3/EQM_R-3/APEQM_R-1"


def air_snapshot():
    """A single-band (B3) dual-pol AIR radio under RUNTIME_VIEW-3, where the
    plan antenna ports (ANTL_A-7/8) differ from the runtime rp1Name ports
    (antenna1a/1b), linked by ANTL_R.configDN. No VSWR MO — only ANTL_M."""
    ant = NP + "/EQM_A-1/APEQM_A-1/RMOD_A-2/ANTL_A-%d"
    return FakeSnap({
        "RMOD_R": [(RV3 + "/RMOD_R-2", {"serialNumber": "FX1",
                                        "productName": "FXED",
                                        "configDN": NP +
                                        "/EQM_A-1/APEQM_A-1/RMOD_A-2"})],
        "RMOD_L": [(EQ + "/RMOD_L-1", {"serialNumber": "FX1"})],
        "FFU_L": [(EQ + "/RMOD_L-1/RU_L-1/FF_L-%d/FFU_L-1" % i,
                   {"centerFrequencyDownlink": 1842e6}) for i in (1, 2)],
        # thresholds on ANTL_R, live VSWR on its ANTL_M child
        "ANTL_R": [(RV3 + "/RMOD_R-2/ANTL_R-1",
                    {"vswrMinorThreshold": 15, "vswrMajorThreshold": 17,
                     "configDN": ant % 7}),
                   (RV3 + "/RMOD_R-2/ANTL_R-2",
                    {"vswrMinorThreshold": 15, "vswrMajorThreshold": 17,
                     "configDN": ant % 8})],
        "ANTL_M": [(RV3 + "/RMOD_R-2/ANTL_R-1/ANTL_M-1",
                    {"rp1Name": "antenna1a", "vswr": 11.0}),
                   (RV3 + "/RMOD_R-2/ANTL_R-2/ANTL_M-1",
                    {"rp1Name": "antenna1b", "vswr": 15.0})],
        "LNCEL_A": [(NP + "/LNBTS_A-9/LNCEL_A-3", {"cellName": "SITE-F-03"})],
        "LNCEL_FDD_A": [(NP + "/LNBTS_A-9/LNCEL_A-3/LNCEL_FDD_A-1",
                         {"earfcnDL": 1650})],          # B3
        "CHANNEL_A": [
            (MAP + "/LCELL_A-3/CHANNELGROUP_A-1/CHANNEL_A-2",
             {"antlDN": ant % 7, "direction": "RX"}),
            (MAP + "/LCELL_A-3/CHANNELGROUP_A-1/CHANNEL_A-3",
             {"antlDN": ant % 8, "direction": "RX"}),
        ],
        "RTWP_MEASUREMENT": [
            (MC + "/LNBTS_M-1/CELL_M-3/CHANNEL_GROUP_M-1/CHANNEL_M-2/"
                  "RTWP_MEASUREMENT-1", {"rtwpValue": -910.0}),
            (MC + "/LNBTS_M-1/CELL_M-3/CHANNEL_GROUP_M-1/CHANNEL_M-3/"
                  "RTWP_MEASUREMENT-1", {"rtwpValue": -960.0}),
        ],
    })


class AirRadioTests(unittest.TestCase):
    def setUp(self):
        self.rows = nr.radio_rows(air_snapshot())

    def test_radio_found_under_runtime_view_3(self):
        self.assertEqual({r["Radio"] for r in self.rows}, {"RMOD-2"})
        self.assertEqual({r["Product"] for r in self.rows}, {"FXED"})

    def test_vswr_from_antl_m_both_branches(self):
        self.assertEqual(sorted(r["VSWR"] for r in self.rows), [1.1, 1.5])
        maj = next(r for r in self.rows if r["VSWR"] == 1.5)
        self.assertEqual(maj["VSWR status"], "Minor")   # 1.5 == minor limit

    def test_single_band_both_pipes_b3(self):
        self.assertEqual({r["Band"] for r in self.rows}, {"B3"})

    def test_cells_attach_through_configdn_port_map(self):
        # plan ports 7/8 map to runtime antenna1a/1b, RTWP attaches per branch
        rx = sorted(r["RX (dBm)"] for r in self.rows if r["Measure"] == "RTWP")
        self.assertEqual(rx, [-96.0, -91.0])
        self.assertTrue(all(r["Cell"] == "SITE-F-03" for r in self.rows))


if __name__ == "__main__":
    unittest.main()
