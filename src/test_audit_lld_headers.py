import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from audit.lld_audit import audit_lld, _sheet_rows, _canonical_header


class LldHeaderTests(unittest.TestCase):
    def test_old_and_v53_port_headers(self):
        for header in ("BB RI Port", "BB RI Port / \nFor Cascade - Radio DATA Port",
                       "Baseband CPRI Port (new allocation)", "BB_RI_PORT / revised notes"):
            with self.subTest(header=header), tempfile.TemporaryDirectory() as folder:
                book = Workbook()
                tx = book.active
                tx.title = "Tx connectivity"
                tx.append(["PLA ID", "BB-1"])
                tx.append(["MIN823", "RP6655"])
                cpri = book.create_sheet("CPRI connectivity")
                cpri.append([1, 2, 3, 4])
                cpri.append(["PLA ID", "BBID", header, "Radio DATA Port"])
                cpri.append(["MIN823", "BB1", "A", "DATA_1"])
                path = Path(folder) / "lld.xlsx"
                book.save(path)
                book.close()
                records = {
                    "ManagedElement=MIN823_GINGOOB01,FieldReplaceableUnit=BB-1":
                        {"productName": "RAN Processor 6655"},
                    "ManagedElement=MIN823_GINGOOB01,RiLink=1": {
                        "riPortRef1": "FieldReplaceableUnit=BB-1,RiPort=A",
                        "riPortRef2": "FieldReplaceableUnit=B28_RRU1,RiPort=DATA_1"},
                }
                rows = audit_lld(str(path), "MIN823_GINGOOB01", records)
                links = [r for r in rows if r.bb_port_lld == "A"]
                self.assertEqual(len(links), 1)
                self.assertEqual(links[0].status, "Match")

    def test_reordered_renamed_and_shifted_headers(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "CPRI connectivity"
        sheet.append(["Revision notes"])
        sheet.append([])
        sheet.append(["RRU DATA PORT", "Site_ID", "BASEBAND ID", "BB CPRI Port - notes"])
        sheet.append(["DATA_1", "MIN823", "BB1", "A"])
        col, rows = _sheet_rows(book, sheet.title, 2)
        self.assertEqual(col, {"Radio DATA Port": 0, "PLA ID": 1, "BBID": 2, "BB RI Port": 3})
        self.assertEqual(rows[0][col["BB RI Port"]], "A")
        book.close()

    def test_ambiguous_columns_rejected_and_unknown_not_guessed(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["PLA ID", "BBID", "BB RI Port", "BB CPRI Port"])
        with self.assertRaisesRegex(ValueError, "ambiguous columns"):
            _sheet_rows(book, sheet.title, 1)
        self.assertEqual(_canonical_header("Radio power port"), "Radio power port")
        self.assertEqual(_canonical_header("BB RI Port / Radio DATA Port"), "BB RI Port")
        book.close()
