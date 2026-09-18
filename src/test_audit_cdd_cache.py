import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from audit import cdd_reader


class CddCacheTests(unittest.TestCase):
    def test_full_rows_reused_across_nodes_and_header_configs(self):
        with tempfile.TemporaryDirectory() as folder:
            path = str(Path(folder) / "cdd.xlsx")
            book = Workbook()
            book.active.title = "CDD"
            book.active.append(["Revision"])
            book.active.append(["Node", "Value"])
            book.active.append(["BB1", 10])
            book.active.append(["BB2", 20])
            book.save(path)
            book.close()
            with patch.object(cdd_reader, "load_workbook", wraps=cdd_reader.load_workbook) as opened:
                with cdd_reader.CddReadCache() as cache:
                    first = cdd_reader._get_sheet(path, "CDD", 1, cache.sheets, cache.workbooks, "Node")
                    second = cdd_reader._get_sheet(path, "CDD", 2, cache.sheets, cache.workbooks, "Node")
                    self.assertEqual(first, second)
                    self.assertEqual(first[1], [("BB1", 10), ("BB2", 20)])
                    self.assertEqual(opened.call_count, 1)
                    self.assertEqual(sum(k[-1] == "complete-rows" for k in cache.sheets), 1)
                self.assertEqual(cache.workbooks, {})
                self.assertEqual(cache.sheets, {})
                # A subsequent audit opens fresh data, not a stale global cache.
                with cdd_reader.CddReadCache() as cache:
                    cdd_reader._get_sheet(path, "CDD", 2, cache.sheets, cache.workbooks, "Node")
                self.assertEqual(opened.call_count, 2)

    def test_cleanup_on_exception(self):
        from unittest.mock import Mock
        cache = cdd_reader.CddReadCache()
        book = Mock()
        with self.assertRaises(RuntimeError):
            with cache:
                cache.workbooks["file"] = book
                raise RuntimeError("cancelled")
        book.close.assert_called_once()
        self.assertFalse(cache.workbooks)
