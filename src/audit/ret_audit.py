"""
ret_audit.py — RET before/after audit for a Nokia → Ericsson swap.

Before: the Nokia site's IM snapshot (.ims2), read with ``ims2_reader``.
After:  an Ericsson moshell log of

    lhgetc AntennaUnitGroup=.*,AntennaNearUnit=

(the AntennaNearUnit table followed by the RetSubUnit table, one log may hold
several nodes).

RET subunits are paired in passes, each on the RETs the previous pass left:

  1. same device   — controller serial + subunit number (burnt into the AISG
                     device, so they survive the swap)
  2. not detected  — the Ericsson RetSubUnit exists but its device was never
                     read (no serial): paired on the configured uniqueId
                     against the device ID Nokia saw, + subunit
  3. replaced      — a different device in the same position: base station ID
                     + subunit, then sector ID + subunit (AISG user data the
                     installer programs into the new RET)

Status per RET — the target is the RET ENABLED on Ericsson:
  OK      — Ericsson RetSubUnit ENABLED and in the same sector as on Nokia
            (a different antenna, a new RET, or one Nokia never had is fine)
  Not OK  — Ericsson DISABLED / absent, or the RET moved to another sector
The remark explains the row: Ericsson state, "mismatch <field>" for every
differing field (unique ID: ``uniqueId`` first, then ``onUnitUniqueId``),
tilt as a note, how a different device was paired, and whether the RET was
already faulty on Nokia.

Sector: Ericsson from the AntennaUnitGroup name (…_S1); Nokia from the radio
the RET controller hangs on (LOGLINK ALD↔RMOD) and the cells that radio
serves (last digit of the local cell ID), else from the RET's own sector ID.

Nokia ↔ Ericsson mapping (runtime view):

  ALD_R  vendorCode+serialNumber  ↔ AntennaNearUnit onUnitUniqueId / uniqueId
  ALD_R  serialNumber             ↔ AntennaNearUnit serialNumber   (pair key)
  ALD_R  productCode / hwVersion  ↔ AntennaNearUnit productNumber / hardwareVersion
  ALD_R  swVersion                ↔ AntennaNearUnit softwareVersion (info only)
  RETU_R subunitNumber            ↔ RetSubUnit subunitNumber        (pair key)
  RETU_R angle / minAngle / maxAngle ↔ electricalAntennaTilt / minTilt / maxTilt
  RETU_R antBearing               ↔ iuantAntennaBearing
  RETU_R antModel / antSerial     ↔ iuantAntennaModelNumber / iuantAntennaSerialNumber
  RETU_R baseStationID / sectorID ↔ iuantBaseStationId / iuantSectorId
  RETU_R installerID              ↔ iuantInstallersId
  RETU_R antBandList[].antOperGain ↔ iuantAntennaOperatingGain
  RETU_R operationalState         ↔ RetSubUnit operationalState
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from datetime import datetime
from typing import Dict, Iterable, List, Optional, Tuple

ERICSSON_COMMAND = "lhgetc AntennaUnitGroup=.*,AntennaNearUnit="

OK, NOT_OK = "OK", "Not OK"
STATUSES = (OK, NOT_OK)
ISSUES = (NOT_OK,)


@dataclass(eq=False)          # identity: two RETs are never "equal"
class Ret:
    side: str                       # "Nokia" / "Ericsson"
    node: str
    mo: str
    serial: str = ""                # controller serial — pair key
    device_id: str = ""             # vendor code + serial as the device reports
    unique_id: str = ""             # configured uniqueId (Nokia: device_id)
    product: str = ""
    hw: str = ""
    sw: str = ""
    subunit: Optional[int] = None   # pair key
    tilt: Optional[int] = None      # 0.1 degree, as AISG reports it
    min_tilt: Optional[int] = None
    max_tilt: Optional[int] = None
    bearing: Optional[int] = None
    ant_model: str = ""
    ant_serial: str = ""
    base_station_id: str = ""
    sector_id: str = ""
    installer: str = ""
    gains: str = ""                 # "158 150 164 160"
    bands: str = ""                 # informational, format differs per vendor
    state: str = ""                 # ENABLED / DISABLED
    sector: str = ""                # physical sector: "S1", "S2", …
    sector_src: str = ""            # where the sector came from
    sector_check: str = ""          # programmed Sector ID vs physical sector
    faults: List[str] = field(default_factory=list)

    @property
    def key(self) -> Tuple[str, Optional[int]]:
        return (self.serial.strip().upper(), self.subunit)

    @property
    def healthy(self) -> bool:
        return not self.faults


# (label, attribute, is tilt) — the fields graded Match / Mismatch.
COMPARE_FIELDS = [
    ("Unique ID", "unique_id", False),
    ("Product", "product", False),
    ("HW version", "hw", False),
    ("Tilt (deg)", "tilt", True),
    ("Min tilt (deg)", "min_tilt", True),
    ("Max tilt (deg)", "max_tilt", True),
    ("Bearing", "bearing", False),
    ("Antenna model", "ant_model", False),
    ("Antenna serial", "ant_serial", False),
    ("Base station ID", "base_station_id", False),
    ("Sector ID", "sector_id", False),
    ("Installer ID", "installer", False),
    ("Gain", "gains", False),
    ("State", "state", False),
]
# Shown side by side but never graded (a swap may legitimately change them).
INFO_FIELDS = [
    ("Sector ID check", "sector_check"),
    ("Controller serial", "serial"),
    ("SW version", "sw"),
    ("Bands", "bands"),
]


def fmt_tilt(v: Optional[int]) -> str:
    return "" if v is None else "%.1f" % (v / 10.0)


def display(ret: Optional[Ret], attr: str, is_tilt: bool = False) -> str:
    if ret is None:
        return ""
    v = getattr(ret, attr)
    if is_tilt:
        return fmt_tilt(v)
    return "" if v is None else str(v)


def _int(v) -> Optional[int]:
    if v is None or v == "":
        return None
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return None


# ── Nokia (.ims2) ─────────────────────────────────────────────────────
def mrbts_from_filename(path: str) -> str:
    m = re.search(r"MRBTS-(\d+)", os.path.basename(path or ""))
    return "MRBTS-%s" % m.group(1) if m else ""


def nokia_rets(snapshot, node: str = "") -> List[Ret]:
    """RETs from a parsed ``ims2_reader.Snapshot``. Uses the runtime view
    (ALD_R/RETU_R, what the BTS actually talks to); falls back to the local
    view (ALD_L/RETU_L) for snapshots that have no runtime view."""
    from ims2_reader import parent
    node = node or mrbts_from_filename(getattr(snapshot, "path", "")) or "Nokia"
    runtime = [(dn, v) for dn, v in snapshot.by_class("RETU_R")
               if "/RUNTIME_VIEW-" in dn]
    if runtime:
        alds = {dn: v for dn, v in snapshot.by_class("ALD_R")
                if "/RUNTIME_VIEW-" in dn}
        rets = [_nokia_runtime(node, dn, v, alds.get(parent(dn), {}))
                for dn, v in runtime]
        by_ald = nokia_ald_sectors(snapshot)
        for (dn, _v), r in zip(runtime, rets):
            sec = by_ald.get(parent(dn))
            if sec:
                r.sector, r.sector_src = sec
        # An ALD without its own radio link (e.g. module M2 of a multi-RET
        # antenna) takes the sector of another module of the same antenna.
        base = lambda r: re.sub(r"-M\d+$", "", r.serial.strip().upper())
        linked = {}
        for r in rets:
            if r.sector and r.serial:
                linked.setdefault(base(r), r)
        for r in rets:
            sib = linked.get(base(r)) if r.serial and not r.sector else None
            if sib:
                r.sector = sib.sector
                r.sector_src = "same antenna as %s" % sib.serial
    else:
        alds = dict(snapshot.by_class("ALD_L"))
        rets = [_nokia_local(node, dn, v, alds.get(parent(dn), {}))
                for dn, v in snapshot.by_class("RETU_L")]
    for r in rets:
        if not r.sector:
            m = re.search(r"SEC-?(\d+)", r.sector_id, re.I)
            if m:
                r.sector, r.sector_src = "S" + m.group(1), "RET sector ID"
        r.sector_check = sector_check(r)
    # An active alarm raised on a RET (or its ALD) marks it faulty before
    # the swap, e.g. 7113 "RET Antenna control failure" on ALD_R-2/RETU_R-3.
    active, _hist = nokia_alarms(snapshot)
    for r in rets:
        ald = r.mo.split("/")[0]
        for a in active:
            obj = a["Alarming object"]
            if obj.endswith("/" + r.mo) or obj.endswith("/" + ald):
                r.faults.append("alarm %s %s" % (a["Alarm"], a["Fault"]
                                                 or a["Name"]))
    return rets


# ── same-site check ───────────────────────────────────────────────────
SITE_MIN_COMMON = 8     # common name core needed to call it the same site


def nokia_bts_name(snapshot) -> str:
    """The Nokia BTS name (MRBTS btsName / LNBTS enbName), else the name in
    the file name (im-snapshot_MRBTS-<id>_<NAME>_SBTS…)."""
    for cls, attr in (("MRBTS_A", "btsName"), ("LNBTS_A", "enbName")):
        names = [str(v.get(attr) or "") for _dn, v in snapshot.by_class(cls)]
        names = [n for n in names if n]
        if names:
            return names[-1]
    m = re.search(r"MRBTS-\d+_([A-Za-z0-9]+)_",
                  os.path.basename(getattr(snapshot, "path", "") or ""))
    return m.group(1) if m else ""


def common_core(a: str, b: str) -> str:
    """Longest common substring of two node names (letters/digits only,
    case-insensitive): 'TCFGACROSRUBTUPISCOTFWLY' and
    'MIN1148_CROSRUBTUPISCOTB01' -> 'CROSRUBTUPISCOT'."""
    a = re.sub(r"[^A-Z0-9]", "", (a or "").upper())
    b = re.sub(r"[^A-Z0-9]", "", (b or "").upper())
    best, best_end = 0, 0
    prev = [0] * (len(b) + 1)
    for i in range(1, len(a) + 1):
        cur = [0] * (len(b) + 1)
        for j in range(1, len(b) + 1):
            if a[i - 1] == b[j - 1]:
                cur[j] = prev[j - 1] + 1
                if cur[j] > best:
                    best, best_end = cur[j], i
        prev = cur
    return a[best_end - best:best_end]


def site_check(nokia_name: str, ericsson_nodes: Iterable[str],
               rows: Optional[List["RetRow"]] = None) -> Dict[str, object]:
    """Do the Nokia BTS and every Ericsson node belong to the same site?

    ``ok`` is True only when each Ericsson node shares a name core of at
    least SITE_MIN_COMMON characters with the Nokia BTS name. RETs paired as
    the same physical device (serial) are reported as extra evidence."""
    nodes = sorted({n for n in ericsson_nodes if n})
    per_node = [(n, common_core(nokia_name, n)) for n in nodes]
    ok = bool(nokia_name) and bool(per_node) and all(
        len(core) >= SITE_MIN_COMMON for _n, core in per_node)
    same_device = sum(1 for r in rows or [] if r.paired_by == "same device")
    if not nokia_name:
        msg = "Nokia BTS name not found in the snapshot"
    elif not per_node:
        msg = "No Ericsson node name found in the log (no moshell prompt)"
    else:
        msg = "; ".join(
            "%s ↔ %s: %s" % (nokia_name, n,
                             "common '%s'" % core if core else "nothing in common")
            for n, core in per_node)
    if same_device:
        msg += " — %d RET(s) are the same physical device on both sides" \
               % same_device
    return {"ok": ok, "nokia": nokia_name, "nodes": per_node,
            "same_device": same_device, "message": msg}


def _alarm_time(v: str) -> str:
    """'20260405192049.281+0800' -> '2026-04-05 19:20:49 +0800'."""
    m = re.match(r"(\d{4})(\d\d)(\d\d)(\d\d)(\d\d)(\d\d)(?:\.\d+)?([+-]\d{4})?$",
                 str(v or ""))
    if not m:
        return str(v or "")
    y, mo, d, h, mi, s, tz = m.groups()
    return "%s-%s-%s %s:%s:%s%s" % (y, mo, d, h, mi, s,
                                    " " + tz if tz else "")


def _alarm_row(dn: str, v: dict, status: str) -> Dict[str, str]:
    text = v.get("alarmText", {}) or {}
    extra = text.get("alarmAdditionalInfo", {}) or {}
    attrs = {a.get("attributeName"): a.get("attributeValue")
             for a in v.get("alarmAttribute", []) or []}
    short = lambda x: re.sub(r"^/MRBTS-1/RAT-1/(RUNTIME_VIEW-1/)?", "",
                             str(x or ""))
    return {
        "Status": status,
        "Alarm": str(v.get("alarmNumber", "")),
        "Name": str(v.get("alarmName", "")),
        "Severity": str(v.get("alarmSeverity", "")),
        "Activity": str(v.get("alarmActivity", "")),
        "Raised": _alarm_time(v.get("observationTime")),
        "Last update": _alarm_time(v.get("lastUpdateTime")),
        "Fault ID": str(v.get("faultId", "")),
        "Fault": str(text.get("faultDescription") or text.get("alarmDetail")
                     or ""),
        "Additional info": str(extra.get("additionalFaultReason", "")),
        "Alarming object": short(v.get("alarmingResourceDN")),
        "Reported by": ", ".join(short(a.get("reportingResourceDN"))
                                 for a in v.get("affectedResources", []) or []),
        "Affected cells": str(attrs.get("affected_cells", "") or ""),
        "Unit": str(attrs.get("unitName", "") or ""),
        "MO": dn.rsplit("/", 1)[-1],
    }


ALARM_COLUMNS = ["Status", "Alarm", "Name", "Severity", "Activity", "Raised",
                 "Last update", "Fault ID", "Fault", "Additional info",
                 "Alarming object", "Reported by", "Affected cells", "Unit",
                 "MO"]


def nokia_alarms(snapshot) -> Tuple[List[dict], List[dict]]:
    """(active, history) alarms of the Nokia BTS.

    Active  = ALARM MOs still present at the end of the snapshot.
    History = every alarm instance written during the snapshot's log, newest
              first, with Status Active or Cleared (its ALARM MO was deleted
              or reused by a later alarm)."""
    decode = lambda p: snapshot.get_payload("ALARM", p)
    active = [_alarm_row(dn, v, "Active")
              for dn, v in snapshot.by_class("ALARM")]
    instances: "Dict[tuple, list]" = {}      # (dn, raised) -> [row data]
    latest: Dict[str, tuple] = {}            # dn -> its latest instance key
    for dn, deleted, payload in getattr(snapshot, "history", []):
        if deleted:
            key = latest.get(dn)
            if key:
                instances[key][1] = "Cleared"
            continue
        v = decode(payload)
        key = (dn, str(v.get("observationTime", "")))
        prev = latest.get(dn)
        if prev and prev != key and instances[prev][1] == "Active":
            instances[prev][1] = "Cleared"   # the MO now holds another alarm
        instances[key] = [v, "Active"]
        latest[dn] = key
    history = [_alarm_row(dn, v, st)
               for (dn, _t), (v, st) in instances.items()]
    by_time = lambda r: r["Raised"]
    active.sort(key=by_time, reverse=True)
    history.sort(key=by_time, reverse=True)
    return active, history


def sector_check(r: Ret) -> str:
    """Cross-check the Sector ID programmed in the RET (AISG user data, e.g.
    'SEC-3_ANT1&2') against the physical sector. '' when either is unknown or
    the sector itself came from that Sector ID."""
    m = re.search(r"SEC-?(\d+)", r.sector_id or "", re.I)
    if not m or not r.sector or r.sector_src == "RET sector ID":
        return ""
    if "S" + m.group(1) == r.sector:
        return "OK"
    return "Sector ID %s but RET is on %s" % (r.sector_id, r.sector)


def nokia_ald_sectors(snapshot) -> Dict[str, Tuple[str, str]]:
    """{ALD_R dn: ("S<n>", source)} — the radio (RMOD) each RET controller is
    linked to (LOGLINK_R) and the cells that radio serves (CHANNEL antlDN);
    the sector is the last digit of those local cell IDs when they agree."""
    rv = "/RUNTIME_VIEW-"
    rmod_of_ald: Dict[str, str] = {}
    for dn, v in snapshot.by_class("LOGLINK_R"):
        if rv not in dn:
            continue
        ends = [str(v.get("firstEndpointDN", "")),
                str(v.get("secondEndpointDN", ""))]
        rmod = next((e for e in ends if "/RMOD_R-" in e), "")
        ald = next((e for e in ends if re.search(r"/ALD_R-\d+$", e)), "")
        if rmod and ald:
            rmod_of_ald[ald] = re.sub(r"(/RMOD_R-\d+).*$", r"\1", rmod)
    if not rmod_of_ald:
        return {}
    # runtime radio -> planned radio (RMOD_A-k) and its plan
    plan_rmod: Dict[str, Tuple[str, str]] = {}
    for dn, v in snapshot.by_class("RMOD_R"):
        m = re.search(r"(/NP-\d+/).*/RMOD_A-(\d+)$", str(v.get("configDN", "")))
        if rv in dn and m:
            plan_rmod[dn] = (m.group(1), m.group(2))
    plans = {np for np, _k in plan_rmod.values()}
    cells: Dict[str, set] = {}
    for dn, v in snapshot.by_class("CHANNEL_A"):
        if not any(np in dn for np in plans):
            continue
        k = re.search(r"/RMOD_A-(\d+)/", str(v.get("antlDN", "")) + "/")
        # LCELL_A-92 (LTE/NR local cell) — not CELLMAPPING_A or LCELC_A (GSM)
        cell = re.search(r"/(?:[A-Z]*CELL|LNCEL)_A-(\d+)/", dn)
        if k and cell:
            cells.setdefault(k.group(1), set()).add(int(cell.group(1)))
    out = {}
    for ald, rmod in rmod_of_ald.items():
        k = plan_rmod.get(rmod, ("", ""))[1]
        digits = {c % 10 for c in cells.get(k, ())} - {0}
        if len(digits) == 1:
            out[ald] = ("S%d" % digits.pop(),
                        "RMOD_R-%s cells" % rmod.rsplit("-", 1)[-1])
    return out


def _short_dn(dn: str) -> str:
    return dn.split("/APEQM_R-1/")[-1].split("/EQM_L-1/")[-1]


def _nokia_runtime(node, dn, v, ald) -> Ret:
    serial = str(ald.get("serialNumber", "") or "")
    device = "%s%s" % (ald.get("vendorCode", "") or "", serial)
    state = str(v.get("operationalState", "") or "").upper()
    r = Ret(
        side="Nokia", node=node, mo=_short_dn(dn), serial=serial,
        device_id=device, unique_id=device,
        product=str(ald.get("productCode", "") or ""),
        hw=str(ald.get("hwVersion", "") or ""),
        sw=str(ald.get("swVersion", "") or ""),
        subunit=_int(v.get("subunitNumber")),
        tilt=_int(v.get("angle")),
        min_tilt=_int(v.get("minAngle")), max_tilt=_int(v.get("maxAngle")),
        bearing=_int(v.get("antBearing")),
        ant_model=str(v.get("antModel", "") or ""),
        ant_serial=str(v.get("antSerial", "") or ""),
        base_station_id=str(v.get("baseStationID", "") or ""),
        sector_id=str(v.get("sectorID", "") or ""),
        installer=str(v.get("installerID", "") or ""),
        gains=" ".join(str(b.get("antOperGain", "")) for b in
                       v.get("antBandList", []) or []),
        bands=" ".join(str(b.get("antFreqBand", "")) for b in
                       v.get("antBandList", []) or []),
        state=state,
    )
    if state != "ENABLED":
        r.faults.append("RETU disabled")
    if r.tilt is None:
        r.faults.append("no tilt reported")
    if str(ald.get("operationalState", "enabled")).lower() != "enabled":
        r.faults.append("ALD disabled")
    if not v.get("configDN"):
        r.faults.append("not in commissioned plan (SCF)")
    return r


def _nokia_local(node, dn, v, ald) -> Ret:
    serial = str(ald.get("serialNumber", "") or "")
    device = "%s%s" % (ald.get("vendorCode", "") or "", serial)
    si = v.get("stateInfo", {}) or {}
    state = str(si.get("operationalState", "") or "").upper()
    tilt = _int(v.get("tiltAngle"))
    if tilt is not None and tilt >= 2550:          # AISG "unknown" sentinel
        tilt = None
    bands = v.get("antennaOperatingBands", []) or []
    r = Ret(
        side="Nokia", node=node, mo=_short_dn(dn), serial=serial,
        device_id=device, unique_id=device,
        product=str(ald.get("productCode", "") or ""),
        hw=str(ald.get("hardwareVersion", "") or ""),
        sw=str(ald.get("softwareVersion", "") or ""),
        subunit=_int(v.get("subunitNumber")), tilt=tilt,
        min_tilt=_int(v.get("minimumSupportedTilt")),
        max_tilt=_int(v.get("maximumSupportedTilt")),
        bearing=_int(v.get("antennaBearing")),
        ant_model=str(v.get("antennaModelNumber", "") or ""),
        ant_serial=str(v.get("antennaSerialNumber", "") or ""),
        base_station_id=str(v.get("baseStationId", "") or ""),
        sector_id=str(v.get("sectorId", "") or ""),
        installer=str(v.get("installerId", "") or ""),
        gains=" ".join(str(_int(b.get("gain")) or "") for b in bands),
        bands=" ".join(str(b.get("band", "")) for b in bands),
        state=state,
    )
    if state != "ENABLED":
        r.faults.append("RETU disabled")
    if r.tilt is None:
        r.faults.append("no tilt reported")
    if si.get("configurationStatus") == "NotConfigured":
        r.faults.append("not configured")
    return r


# ── Ericsson (lhgetc log) ─────────────────────────────────────────────
_PROMPT = re.compile(r"^\s*([A-Za-z0-9_.-]+)>\s*(.*)$")
_ENUM = re.compile(r"^-?\d+\s+\((.*)\)$")
_ARRAY = re.compile(r"^i?\[\d+\]\s*=\s*(.*)$")


def _value(raw: str) -> str:
    """'1 (UNLOCKED)' -> 'UNLOCKED', 'i[1] = 1 (FAILED)' -> 'FAILED',
    'i[4] = 158 150' -> '158 150', 'i[0] =' -> ''."""
    v = raw.strip()
    m = _ARRAY.match(v)
    if m:
        v = m.group(1).strip()
    m = _ENUM.match(v)
    if m:
        v = m.group(1).strip()
    return v


def parse_lhgetc(text: str) -> Dict[str, Dict[str, Dict[str, str]]]:
    """{node: {MO: {attribute: value}}} from moshell ``lhgetc`` output.
    Continuation rows (struct/array members printed on the next line with
    the same MO) are merged into the MO's first row."""
    out: Dict[str, Dict[str, Dict[str, str]]] = {}
    node = ""
    header: List[str] = []
    for line in text.splitlines():
        line = line.rstrip("\r\n")
        m = _PROMPT.match(line)
        if m and ";" not in line:
            node = m.group(1)
            header = []
            continue
        if ";" not in line:
            continue
        parts = [p.strip() for p in line.split(";")]
        if parts[0] == "MO":
            header = parts
            continue
        if not header or "=" not in parts[0]:
            continue
        attrs = out.setdefault(node, {}).setdefault(parts[0], {})
        for name, raw in zip(header[1:], parts[1:]):
            val = _value(raw)
            if val and not attrs.get(name):
                attrs[name] = val
    return out


_SECTOR_RE = re.compile(r"[_-]S(\d+)$", re.I)


def ericsson_sector(mo: str, anu: dict, rsu: dict) -> Tuple[str, str]:
    """Sector from the AntennaUnitGroup name (…_S1), else the group the
    RetSubUnit is reserved by, else the radio it hangs on (…RRU1)."""
    m = re.search(r"AntennaUnitGroup=([^,]+)", mo)
    if m and _SECTOR_RE.search(m.group(1)):
        return "S" + _SECTOR_RE.search(m.group(1)).group(1), "AntennaUnitGroup"
    m = re.search(r"AntennaUnitGroup=([^,]+)", rsu.get("reservedBy", ""))
    if m and _SECTOR_RE.search(m.group(1)):
        return ("S" + _SECTOR_RE.search(m.group(1)).group(1),
                "reservedBy AntennaUnitGroup")
    m = re.search(r"RRU_?(\d+)\b", anu.get("rfPortRef", ""), re.I)
    if m:
        return "S" + m.group(1), "rfPortRef radio"
    return "", ""


def ericsson_rets(text: str) -> List[Ret]:
    rets: List[Ret] = []
    for node, mos in parse_lhgetc(text).items():
        for mo, a in mos.items():
            if ",RetSubUnit=" not in mo:
                continue
            anu = mos.get(mo.split(",RetSubUnit=")[0], {})
            state = a.get("operationalState", "").upper()
            r = Ret(
                side="Ericsson", node=node, mo=mo,
                serial=anu.get("serialNumber", ""),
                device_id=anu.get("onUnitUniqueId", ""),
                unique_id=anu.get("uniqueId", ""),
                product=anu.get("productNumber", ""),
                hw=anu.get("hardwareVersion", ""),
                sw=anu.get("softwareVersion", ""),
                subunit=_int(a.get("subunitNumber") or a.get("retSubUnitId")),
                tilt=_int(a.get("electricalAntennaTilt")),
                min_tilt=_int(a.get("minTilt")), max_tilt=_int(a.get("maxTilt")),
                bearing=_int(a.get("iuantAntennaBearing")),
                ant_model=a.get("iuantAntennaModelNumber", ""),
                ant_serial=a.get("iuantAntennaSerialNumber", ""),
                base_station_id=a.get("iuantBaseStationId", ""),
                sector_id=a.get("iuantSectorId", ""),
                installer=a.get("iuantInstallersId", ""),
                gains=a.get("iuantAntennaOperatingGain", ""),
                bands=a.get("iuantAntennaOperatingBand", ""),
                state=state,
            )
            r.sector, r.sector_src = ericsson_sector(mo, anu, a)
            r.sector_check = sector_check(r)
            if not r.serial:
                r.faults.append("device not detected")
            if state != "ENABLED":
                r.faults.append("RetSubUnit %s" % (state or "state unknown"))
            if anu.get("operationalState", "").upper() == "DISABLED":
                r.faults.append("AntennaNearUnit DISABLED")
            if "FAILED" in a.get("availabilityStatus", "").upper():
                r.faults.append("availability FAILED")
            if "FAILED" in a.get("calibrationStatus", "").upper():
                r.faults.append("calibration FAILED")
            rets.append(r)
    return rets


# ── compare ───────────────────────────────────────────────────────────
@dataclass
class RetRow:
    status: str
    nokia: Optional[Ret]
    ericsson: Optional[Ret]
    diffs: List[str] = field(default_factory=list)   # labels that differ
    remark: str = ""
    paired_by: str = ""             # which pass paired the two sides

    @property
    def any(self) -> Ret:
        return self.nokia or self.ericsson

    @property
    def is_issue(self) -> bool:
        return self.status in ISSUES


def _norm(ret: Ret, attr: str):
    v = getattr(ret, attr)
    return v.strip() if isinstance(v, str) else v


def unique_id_matches(n: Ret, e: Ret) -> bool:
    """The Nokia device ID against Ericsson ``uniqueId`` first, then
    ``onUnitUniqueId`` (what the device itself reports)."""
    want = n.unique_id.strip().upper()
    return bool(want) and want in (e.unique_id.strip().upper(),
                                   e.device_id.strip().upper())


def ericsson_unique_id(n: Optional[Ret], e: Optional[Ret]) -> str:
    """The Ericsson unique ID to show next to Nokia's: ``uniqueId`` when it
    matches (or there is nothing to match), else ``onUnitUniqueId`` — the
    value the comparison fell back to."""
    if e is None:
        return ""
    if n is None or not e.device_id:
        return e.unique_id or e.device_id
    want = n.unique_id.strip().upper()
    if e.unique_id.strip().upper() == want:
        return e.unique_id
    return e.device_id


def _diffs(n: Ret, e: Ret) -> List[str]:
    out = []
    for label, attr, _t in COMPARE_FIELDS:
        if attr == "unique_id":
            if not unique_id_matches(n, e):
                out.append(label)
        elif _norm(n, attr) != _norm(e, attr):
            out.append(label)
    return out


def _faults(r: Ret) -> str:
    return ", ".join(r.faults)


def _pair_unique(nokia: List[Ret], ericsson: List[Ret], nkey, ekey):
    """Pair on a key only where it is unique on both sides (a shared sector ID
    across two controllers must not pair the wrong RETs)."""
    def index(rets, keyf):
        out: Dict[tuple, List[Ret]] = {}
        for r in rets:
            k = keyf(r)
            if k and all(x not in ("", None) for x in k):
                out.setdefault(k, []).append(r)
        return {k: v[0] for k, v in out.items() if len(v) == 1}
    ni, ei = index(nokia, nkey), index(ericsson, ekey)
    return [(ni[k], ei[k]) for k in ni if k in ei]


def sector_counts(nokia: Iterable[Ret], ericsson: Iterable[Ret]):
    """(sorted Nokia sectors, sorted Ericsson sectors) that have a RET."""
    ns = sorted({r.sector for r in nokia if r.sector})
    es = sorted({r.sector for r in ericsson if r.sector})
    return ns, es


def _enabled(e: Optional[Ret]) -> bool:
    return e is not None and e.state == "ENABLED"


def _ericsson_state(e: Ret) -> str:
    detail = [f for f in e.faults if not f.startswith("RetSubUnit ")]
    return "Ericsson %s%s" % (e.state or "state unknown",
                              " (%s)" % ", ".join(detail) if detail else "")


def compare(nokia: Iterable[Ret], ericsson: Iterable[Ret]) -> List[RetRow]:
    """Pair the RETs of both sides and grade them against the target: the
    RET is ENABLED on Ericsson in the same sector as on Nokia.

    OK      — Ericsson RetSubUnit ENABLED and the sector is unchanged (a
              different antenna, a new RET or one Nokia never had is fine)
    Not OK  — Ericsson DISABLED / absent, or the RET moved to another sector

    Every other difference (unique ID, antenna, tilt, …) only explains the
    row in the remark; the differing cells are coloured in the report."""
    nokia, ericsson = list(nokia), list(ericsson)
    rows: List[RetRow] = []
    left_n, left_e = list(nokia), list(ericsson)

    def take(pairs, how):
        for n, e in pairs:
            if n in left_n and e in left_e:
                left_n.remove(n)
                left_e.remove(e)
                rows.append(_grade(n, e, how))

    up = lambda s: s.strip().upper()
    # 1. same device
    take(_pair_unique(left_n, left_e, lambda r: r.key,
                      lambda r: r.key if r.serial else None), "same device")
    # 2. Ericsson MO whose device was never read: configured uniqueId
    take(_pair_unique(left_n, [e for e in left_e if not e.serial],
                      lambda r: (up(r.device_id), r.subunit),
                      lambda r: (up(r.unique_id), r.subunit)),
         "configured uniqueId")
    # 3. replaced: same position (AISG user data programmed on the new RET)
    detected = lambda: [e for e in left_e if e.serial]
    take(_pair_unique(left_n, detected(),
                      lambda r: (up(r.base_station_id), r.subunit),
                      lambda r: (up(r.base_station_id), r.subunit)),
         "base station ID")
    take(_pair_unique(left_n, detected(),
                      lambda r: (up(r.sector_id), r.subunit),
                      lambda r: (up(r.sector_id), r.subunit)),
         "sector ID")
    # 4. one side only
    for n in left_n:
        notes = ["not found on Ericsson"]
        if n.faults:
            notes.append("Nokia already: %s" % _faults(n))
        rows.append(RetRow(NOT_OK, n, None, remark="; ".join(notes)))
    for e in left_e:
        notes = ["new RET (not on Nokia)"]
        if not _enabled(e):
            notes.insert(0, _ericsson_state(e))
        rows.append(RetRow(OK if _enabled(e) else NOT_OK, None, e,
                           remark="; ".join(notes)))
    # Sector ID programmed in the RET vs its physical sector — a note.
    for row in rows:
        warn = ["%s: %s" % (r.side, r.sector_check)
                for r in (row.nokia, row.ericsson)
                if r is not None and r.sector_check not in ("", "OK")]
        if warn:
            row.remark = "; ".join(filter(None, [row.remark] + warn))
    rows.sort(key=lambda r: (r.any.sector or "S~", r.any.serial,
                             r.any.subunit or 0))
    return rows


def _grade(n: Ret, e: Ret, how: str) -> RetRow:
    diffs = _diffs(n, e)
    moved = bool(n.sector and e.sector and n.sector != e.sector)
    status = OK if _enabled(e) and not moved else NOT_OK
    notes = []
    if not _enabled(e):
        notes.append(_ericsson_state(e))
    if moved:
        notes.append("sector changed %s → %s" % (n.sector, e.sector))
    elif not (n.sector and e.sector):
        notes.append("sector not verified")
    mism = [d for d in diffs if d not in ("Tilt (deg)", "State")]
    if mism:
        notes.append("mismatch %s" % ", ".join(mism))
    if "Tilt (deg)" in diffs:
        notes.append("tilt %s° → %s° (note)" % (fmt_tilt(n.tilt) or "-",
                                                 fmt_tilt(e.tilt) or "-"))
    if how == "configured uniqueId":
        notes.append("device not read on Ericsson, paired by configured "
                     "uniqueId")
    elif how != "same device":
        notes.append("different device, paired by %s" % how)
    if n.faults:
        notes.append("Nokia already: %s" % _faults(n))
    return RetRow(status, n, e, diffs, "; ".join(notes), how)


def summarize(rows: List[RetRow]) -> Dict[str, int]:
    out = {s: 0 for s in STATUSES}
    for r in rows:
        out[r.status] = out.get(r.status, 0) + 1
    return out


def run(ims2_path: str, log_paths: List[str], node: str = ""):
    """Read both sides and compare. Returns (rows, nokia_rets, ericsson_rets,
    (active_alarms, alarm_history), site_check)."""
    from ims2_reader import Snapshot
    snap = Snapshot(ims2_path)
    nk = nokia_rets(snap, node=node)
    er: List[Ret] = []
    for p in log_paths:
        with open(p, "r", encoding="utf-8", errors="replace") as fh:
            er.extend(ericsson_rets(fh.read()))
    rows = compare(nk, er)
    return (rows, nk, er, nokia_alarms(snap),
            site_check(nokia_bts_name(snap), [r.node for r in er], rows))


# ── Excel ─────────────────────────────────────────────────────────────
def write_excel(rows: List[RetRow], nokia: List[Ret], ericsson: List[Ret],
                path: str, meta: Dict[str, str], alarms=None) -> str:
    """``alarms`` = (active, history) from :func:`nokia_alarms` — adds the
    "Nokia Alarms" sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    fills = {
        OK: PatternFill("solid", fgColor="C6EFCE"),
        NOT_OK: PatternFill("solid", fgColor="FFC7CE"),
    }
    diff_fill = PatternFill("solid", fgColor="FFC7CE")
    head_fill = PatternFill("solid", fgColor="4472C4")
    nok_fill = PatternFill("solid", fgColor="7F6000")
    eri_fill = PatternFill("solid", fgColor="1F4E78")
    head_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    left = Alignment(horizontal="left", vertical="center")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = Workbook()

    # Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "RET Audit — Nokia (before) vs Ericsson (after)"
    ws["A1"].font = Font(bold=True, size=16, color="1F3864")
    r = 3
    for k, v in meta.items():
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, v)
        if k == "Same-site check":
            ws.cell(r, 2).fill = fills[OK if str(v).startswith("OK")
                                      else NOT_OK]
            ws.cell(r, 2).font = Font(bold=True)
        r += 1
    r += 1
    for c, h in enumerate(("Status", "Count", "Meaning"), 1):
        cell = ws.cell(r, c, h)
        cell.fill, cell.font = head_fill, head_font
    meaning = {
        OK: "ENABLED on Ericsson in the same sector as on Nokia",
        NOT_OK: "DISABLED or absent on Ericsson, or moved to another sector "
                "— see Remark",
    }
    for status, n in summarize(rows).items():
        r += 1
        ws.cell(r, 1, status).fill = fills[status]
        ws.cell(r, 2, n)
        ws.cell(r, 3, meaning[status])
    ns, es = sector_counts(nokia, ericsson)
    r += 2
    ws.cell(r, 1, "Sectors with RET").font = Font(bold=True)
    ws.cell(r, 2, "Nokia %d (%s)  |  Ericsson %d (%s)%s" % (
        len(ns), " ".join(ns) or "-", len(es), " ".join(es) or "-",
        "" if len(ns) == len(es) else "  — COUNT CHANGED"))
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 60

    # RET Audit (side by side)
    ws = wb.create_sheet("RET Audit")
    fixed = ["Status", "Remark", "Paired by", "Sector\nNokia",
             "Sector\nEricsson", "Subunit", "Nokia MO", "Ericsson node",
             "Ericsson MO"]
    pairs = [(label, attr, t, True) for label, attr, t in COMPARE_FIELDS]
    pairs += [(label, attr, False, False) for label, attr in INFO_FIELDS]
    headers = list(fixed)
    for label, _a, _t, _g in pairs:
        headers += ["%s\nNokia" % label, "%s\nEricsson" % label]
    headers += ["Faults Nokia", "Faults Ericsson"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font, cell.alignment, cell.border = head_font, center, border
        cell.fill = (nok_fill if h.endswith("Nokia") else
                     eri_fill if h.endswith("Ericsson") else head_fill)
    for i, row in enumerate(rows, 2):
        a = row.any
        vals = [row.status, row.remark, row.paired_by,
                row.nokia.sector if row.nokia else "",
                row.ericsson.sector if row.ericsson else "",
                "" if a.subunit is None else a.subunit,
                row.nokia.mo if row.nokia else "",
                row.ericsson.node if row.ericsson else "",
                row.ericsson.mo if row.ericsson else ""]
        for label, attr, is_tilt, _g in pairs:
            vals += [display(row.nokia, attr, is_tilt),
                     ericsson_unique_id(row.nokia, row.ericsson)
                     if attr == "unique_id"
                     else display(row.ericsson, attr, is_tilt)]
        vals += [", ".join(row.nokia.faults) if row.nokia else "",
                 ", ".join(row.ericsson.faults) if row.ericsson else ""]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(i, c, v)
            cell.border, cell.alignment = border, left
        ws.cell(i, 1).fill = fills.get(row.status, fills[NOT_OK])
        col = len(fixed) + 1
        for label, _a, _t, graded in pairs:
            if graded and label in row.diffs:
                ws.cell(i, col).fill = diff_fill
                ws.cell(i, col + 1).fill = diff_fill
            col += 2
    ws.freeze_panes = "B2"         # header row + the Status column
    ws.row_dimensions[1].height = 32
    wide = {"Remark": 60, "Nokia MO": 26, "Ericsson MO": 44,
            "Ericsson node": 26, "Paired by": 18}
    for c, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(c)].width = wide.get(
            h, 14 if c > len(fixed) else 10)
    ws.auto_filter.ref = ws.dimensions

    # Raw per side
    for title, rets in (("Nokia RET", nokia), ("Ericsson RET", ericsson)):
        ws = wb.create_sheet(title)
        cols = [("Node", "node", False), ("MO", "mo", False),
                ("Controller serial", "serial", False),
                ("Device ID", "device_id", False),
                ("Unique ID", "unique_id", False),
                ("Sector", "sector", False),
                ("Sector source", "sector_src", False),
                ("Subunit", "subunit", False)]
        cols += [(l, a, t) for l, a, t in COMPARE_FIELDS
                 if a not in ("unique_id",)]
        cols += [(l, a, False) for l, a in INFO_FIELDS if a != "serial"]
        for c, (label, _a, _t) in enumerate(cols + [("Faults", "", False)], 1):
            cell = ws.cell(1, c, label)
            cell.fill, cell.font, cell.border = head_fill, head_font, border
        for i, ret in enumerate(rets, 2):
            for c, (_l, attr, is_tilt) in enumerate(cols, 1):
                ws.cell(i, c, display(ret, attr, is_tilt)).border = border
            ws.cell(i, len(cols) + 1, ", ".join(ret.faults)).border = border
        for c in range(1, len(cols) + 2):
            ws.column_dimensions[get_column_letter(c)].width = 18
        ws.column_dimensions["B"].width = 44
        ws.auto_filter.ref = ws.dimensions

    if alarms is not None:
        _write_alarm_sheet(wb.create_sheet("Nokia Alarms"), *alarms)

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)
    return path


def _write_alarm_sheet(ws, active: List[dict], history: List[dict]):
    """Active alarms on top, the alarm history below."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="4472C4")
    head_font = Font(bold=True, color="FFFFFF")
    sev_fill = {
        "Critical": PatternFill("solid", fgColor="FF9999"),
        "Major": PatternFill("solid", fgColor="F8CBAD"),
        "Minor": PatternFill("solid", fgColor="FFEB9C"),
        "Warning": PatternFill("solid", fgColor="DDEBF7"),
    }
    cleared_font = Font(color="808080")
    left = Alignment(horizontal="left", vertical="center")

    r = 1

    def section(title, rows, colour):
        nonlocal r
        cell = ws.cell(r, 1, "%s (%d)" % (title, len(rows)))
        cell.font = Font(bold=True, size=13, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=colour)
        ws.merge_cells(start_row=r, start_column=1, end_row=r,
                       end_column=len(ALARM_COLUMNS))
        r += 1
        for c, h in enumerate(ALARM_COLUMNS, 1):
            hc = ws.cell(r, c, h)
            hc.fill, hc.font, hc.border = head_fill, head_font, border
        r += 1
        if not rows:
            ws.cell(r, 1, "none").font = cleared_font
            r += 1
        for row in rows:
            for c, h in enumerate(ALARM_COLUMNS, 1):
                cell = ws.cell(r, c, row.get(h, ""))
                cell.border, cell.alignment = border, left
                if row.get("Status") == "Cleared":
                    cell.font = cleared_font
            fill = sev_fill.get(row.get("Severity"))
            if fill:
                ws.cell(r, ALARM_COLUMNS.index("Severity") + 1).fill = fill
            r += 1

    section("ACTIVE ALARMS", active, "C00000")
    r += 1
    section("ALARM HISTORY (from the snapshot log, newest first)", history,
            "595959")
    widths = {"Name": 38, "Fault": 40, "Additional info": 60,
              "Alarming object": 40, "Reported by": 50, "Raised": 26,
              "Last update": 26}
    for c, h in enumerate(ALARM_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(h, 12)


def default_meta(ims2_path: str, log_paths: List[str], site: str,
                 check: Optional[Dict[str, object]] = None) -> Dict[str, str]:
    meta = {
        "Site": site,
        "Nokia snapshot (before)": os.path.basename(ims2_path),
        "Ericsson log(s) (after)": ", ".join(os.path.basename(p)
                                             for p in log_paths),
        "Ericsson command": ERICSSON_COMMAND,
        "Generated": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }
    if check is not None:
        meta["Nokia BTS name"] = str(check["nokia"] or "-")
        meta["Ericsson node(s)"] = ", ".join(n for n, _c in check["nodes"]) or "-"
        meta["Same-site check"] = "%s — %s" % (
            "OK" if check["ok"] else "CHECK: names do not match",
            check["message"])
    return meta
